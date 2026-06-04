import io
import os
import logging
from datetime import datetime, date, timedelta
from typing import Optional, List, Dict, Any, Tuple
from decimal import Decimal

from sqlalchemy import select, and_, func, case, extract
from sqlalchemy.ext.asyncio import AsyncSession

from .models import (
    Employee, Department, Shift, Schedule, PunchRecord, PunchType,
    AnomalyRecord, AnomalyType, AnomalyStatus, Holiday,
    MonthlyAttendanceSummary, DashboardMetric, ScheduleStatus,
)
from .config import get_config, SystemConfig

logger = logging.getLogger(__name__)


class MonthlyCalculator:
    def __init__(self, config: SystemConfig = None):
        self.config = config or get_config()

    async def calculate_employee_summary(
        self, employee_id: int, year_month: str, session: AsyncSession,
    ) -> MonthlyAttendanceSummary:
        year, month = int(year_month[:4]), int(year_month[5:7])
        start_date = date(year, month, 1)
        if month == 12:
            end_date = date(year + 1, 1, 1) - timedelta(days=1)
        else:
            end_date = date(year, month + 1, 1) - timedelta(days=1)

        employee = await session.get(Employee, employee_id)
        if not employee:
            raise ValueError(f"员工不存在: {employee_id}")

        stmt_schedule = select(func.count()).where(
            and_(
                Schedule.employee_id == employee_id,
                Schedule.schedule_date >= start_date,
                Schedule.schedule_date <= end_date,
                Schedule.status == ScheduleStatus.PUBLISHED,
            )
        )
        result = await session.execute(stmt_schedule)
        total_work_days = result.scalar() or 0

        stmt_punch_days = select(func.count(func.distinct(PunchRecord.punch_date))).where(
            and_(
                PunchRecord.employee_id == employee_id,
                PunchRecord.punch_date >= start_date,
                PunchRecord.punch_date <= end_date,
                PunchRecord.is_valid == True,
            )
        )
        result = await session.execute(stmt_punch_days)
        actual_work_days = result.scalar() or 0

        stmt_anomaly = select(
            AnomalyRecord.anomaly_type,
            func.count().label("cnt"),
        ).where(
            and_(
                AnomalyRecord.employee_id == employee_id,
                AnomalyRecord.anomaly_date >= start_date,
                AnomalyRecord.anomaly_date <= end_date,
            )
        ).group_by(AnomalyRecord.anomaly_type)
        result = await session.execute(stmt_anomaly)

        late_count = 0
        early_leave_count = 0
        absent_count = 0
        missing_punch_count = 0
        for row in result:
            if row.anomaly_type == AnomalyType.LATE:
                late_count = row.cnt
            elif row.anomaly_type == AnomalyType.EARLY_LEAVE:
                early_leave_count = row.cnt
            elif row.anomaly_type == AnomalyType.ABSENT:
                absent_count = row.cnt
            elif row.anomaly_type in (AnomalyType.MISSING_PUNCH_IN, AnomalyType.MISSING_PUNCH_OUT):
                missing_punch_count += row.cnt

        stmt_anomaly_active = select(
            AnomalyRecord.anomaly_type,
            AnomalyRecord.status,
            func.count().label("cnt"),
        ).where(
            and_(
                AnomalyRecord.employee_id == employee_id,
                AnomalyRecord.anomaly_date >= start_date,
                AnomalyRecord.anomaly_date <= end_date,
            )
        ).group_by(AnomalyRecord.anomaly_type, AnomalyRecord.status)
        result_active = await session.execute(stmt_anomaly_active)

        corrected_late = 0
        corrected_early = 0
        corrected_absent = 0
        corrected_missing = 0
        for row in result_active:
            if row.status in (AnomalyStatus.AUTO_CORRECTED, AnomalyStatus.APPROVED):
                if row.anomaly_type == AnomalyType.LATE:
                    corrected_late = row.cnt
                elif row.anomaly_type == AnomalyType.EARLY_LEAVE:
                    corrected_early = row.cnt
                elif row.anomaly_type == AnomalyType.ABSENT:
                    corrected_absent = row.cnt
                elif row.anomaly_type in (AnomalyType.MISSING_PUNCH_IN, AnomalyType.MISSING_PUNCH_OUT):
                    corrected_missing += row.cnt

        effective_late = late_count - corrected_late
        effective_early = early_leave_count - corrected_early
        effective_absent = absent_count - corrected_absent
        effective_missing = missing_punch_count - corrected_missing

        total_work_hours = Decimal("0")
        overtime_weekday = Decimal("0")
        overtime_weekend = Decimal("0")
        overtime_holiday = Decimal("0")

        stmt_punch = select(PunchRecord).where(
            and_(
                PunchRecord.employee_id == employee_id,
                PunchRecord.punch_date >= start_date,
                PunchRecord.punch_date <= end_date,
                PunchRecord.is_valid == True,
            )
        ).order_by(PunchRecord.punch_date, PunchRecord.punch_time)
        result = await session.execute(stmt_punch)
        punches = list(result.scalars().all())

        from collections import defaultdict
        daily_punches = defaultdict(list)
        for p in punches:
            daily_punches[p.punch_date].append(p)

        for punch_date, day_punches in daily_punches.items():
            clock_ins = [p for p in day_punches if p.punch_type == PunchType.CLOCK_IN]
            clock_outs = [p for p in day_punches if p.punch_type == PunchType.CLOCK_OUT]

            if clock_ins and clock_outs:
                earliest_in = min(p.punch_time for p in clock_ins)
                latest_out = max(p.punch_time for p in clock_outs)
                hours = Decimal(str(round((latest_out - earliest_in).total_seconds() / 3600, 2)))

                total_work_hours += hours

                standard_hours = Decimal("8")
                if hours > standard_hours:
                    ot_hours = hours - standard_hours
                    is_weekend = punch_date.weekday() >= 5
                    is_holiday_day = False

                    stmt_hol = select(Holiday).where(Holiday.holiday_date == punch_date)
                    hol_result = await session.execute(stmt_hol)
                    hol = hol_result.scalar_one_or_none()
                    if hol and hol.is_day_off:
                        is_holiday_day = True

                    if is_holiday_day:
                        overtime_holiday += ot_hours
                    elif is_weekend:
                        overtime_weekend += ot_hours
                    else:
                        overtime_weekday += ot_hours

        salary_cfg = self.config.salary
        salary_deduction = (
            Decimal(str(salary_cfg.late_deduction_per_minute)) * effective_late * 15
            + Decimal(str(salary_cfg.early_leave_deduction_per_minute)) * effective_early * 15
            + Decimal(str(salary_cfg.absent_deduction_per_day)) * effective_absent
            + Decimal(str(salary_cfg.missing_punch_deduction)) * effective_missing
        )

        overtime_allowance = (
            Decimal(str(salary_cfg.overtime_rate_weekday)) * overtime_weekday
            + Decimal(str(salary_cfg.overtime_rate_weekend)) * overtime_weekend
            + Decimal(str(salary_cfg.overtime_rate_holiday)) * overtime_holiday
        ) * (employee.hourly_rate if employee.hourly_rate else Decimal("1"))

        net_adjustment = overtime_allowance - salary_deduction
        compliance_rate = Decimal("1.0") if total_work_days == 0 else Decimal(str(
            round(actual_work_days / total_work_days, 4)
        ))

        existing_stmt = select(MonthlyAttendanceSummary).where(
            and_(
                MonthlyAttendanceSummary.employee_id == employee_id,
                MonthlyAttendanceSummary.year_month == year_month,
            )
        )
        existing_result = await session.execute(existing_stmt)
        summary = existing_result.scalar_one_or_none()

        if summary:
            summary.total_work_days = total_work_days
            summary.actual_work_days = actual_work_days
            summary.late_count = effective_late
            summary.early_leave_count = effective_early
            summary.absent_count = effective_absent
            summary.missing_punch_count = effective_missing
            summary.total_work_hours = total_work_hours
            summary.overtime_hours_weekday = overtime_weekday
            summary.overtime_hours_weekend = overtime_weekend
            summary.overtime_hours_holiday = overtime_holiday
            summary.salary_deduction = salary_deduction
            summary.overtime_allowance = overtime_allowance
            summary.net_adjustment = net_adjustment
            summary.compliance_rate = compliance_rate
            summary.generated_at = datetime.now()
        else:
            summary = MonthlyAttendanceSummary(
                employee_id=employee_id,
                department_id=employee.department_id,
                year_month=year_month,
                total_work_days=total_work_days,
                actual_work_days=actual_work_days,
                late_count=effective_late,
                early_leave_count=effective_early,
                absent_count=effective_absent,
                missing_punch_count=effective_missing,
                total_work_hours=total_work_hours,
                overtime_hours_weekday=overtime_weekday,
                overtime_hours_weekend=overtime_weekend,
                overtime_hours_holiday=overtime_holiday,
                salary_deduction=salary_deduction,
                overtime_allowance=overtime_allowance,
                net_adjustment=net_adjustment,
                compliance_rate=compliance_rate,
            )
            session.add(summary)

        await session.flush()
        return summary

    async def calculate_department_summaries(
        self, department_id: int, year_month: str, session: AsyncSession,
    ) -> List[MonthlyAttendanceSummary]:
        stmt = select(Employee).where(
            and_(Employee.department_id == department_id, Employee.is_active == True)
        )
        result = await session.execute(stmt)
        employees = list(result.scalars().all())

        summaries = []
        for emp in employees:
            summary = await self.calculate_employee_summary(emp.id, year_month, session)
            summaries.append(summary)

        logger.info("部门 %d 月度统计完成 [%s]: %d 人", department_id, year_month, len(summaries))
        return summaries

    async def calculate_all_summaries(self, year_month: str, session: AsyncSession) -> List[MonthlyAttendanceSummary]:
        stmt = select(Department).where(Department.is_active == True)
        result = await session.execute(stmt)
        departments = list(result.scalars().all())

        all_summaries = []
        for dept in departments:
            summaries = await self.calculate_department_summaries(dept.id, year_month, session)
            all_summaries.extend(summaries)

        logger.info("全公司月度统计完成 [%s]: %d 人", year_month, len(all_summaries))
        return all_summaries


class ReportGenerator:
    def __init__(self, config: SystemConfig = None):
        self.config = config or get_config()
        self.output_dir = self.config.reports_output_dir
        os.makedirs(self.output_dir, exist_ok=True)
        self.calculator = MonthlyCalculator(self.config)

    async def generate_excel_report(
        self, year_month: str, session: AsyncSession,
        department_id: int = None,
    ) -> str:
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        except ImportError:
            logger.error("缺少 openpyxl 依赖，无法生成Excel报告")
            return ""

        if department_id:
            summaries = await self.calculator.calculate_department_summaries(
                department_id, year_month, session,
            )
        else:
            summaries = await self.calculator.calculate_all_summaries(year_month, session)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"考勤统计{year_month}"

        headers = [
            "工号", "姓名", "部门", "应出勤天数", "实出勤天数",
            "迟到次数", "早退次数", "缺勤天数", "缺卡次数",
            "总工时(h)", "平日加班(h)", "周末加班(h)", "节假日加班(h)",
            "薪资扣减(元)", "加班补贴(元)", "净调整(元)", "达标率",
        ]
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for row_idx, summary in enumerate(summaries, 2):
            emp = await session.get(Employee, summary.employee_id)
            dept = await session.get(Department, summary.department_id)

            values = [
                emp.employee_no if emp else "",
                emp.name if emp else "",
                dept.name if dept else "",
                summary.total_work_days,
                summary.actual_work_days,
                summary.late_count,
                summary.early_leave_count,
                summary.absent_count,
                summary.missing_punch_count,
                float(summary.total_work_hours),
                float(summary.overtime_hours_weekday),
                float(summary.overtime_hours_weekend),
                float(summary.overtime_hours_holiday),
                float(summary.salary_deduction),
                float(summary.overtime_allowance),
                float(summary.net_adjustment),
                f"{float(summary.compliance_rate) * 100:.1f}%",
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.alignment = Alignment(horizontal="center")
                cell.border = thin_border

        for col in range(1, len(headers) + 1):
            max_length = max(
                len(str(ws.cell(row=r, column=col).value or ""))
                for r in range(1, len(summaries) + 2)
            )
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(max_length + 4, 30)

        filename = f"attendance_report_{year_month}"
        if department_id:
            filename += f"_dept{department_id}"
        filename += ".xlsx"
        filepath = os.path.join(self.output_dir, filename)
        wb.save(filepath)

        logger.info("Excel报告已生成: %s", filepath)
        return filepath

    async def generate_pdf_report(
        self, year_month: str, session: AsyncSession,
        department_id: int = None,
    ) -> str:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import mm
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        except ImportError:
            logger.error("缺少 reportlab 依赖，无法生成PDF报告")
            return ""

        if department_id:
            summaries = await self.calculator.calculate_department_summaries(
                department_id, year_month, session,
            )
        else:
            summaries = await self.calculator.calculate_all_summaries(year_month, session)

        filename = f"attendance_report_{year_month}"
        if department_id:
            filename += f"_dept{department_id}"
        filename += ".pdf"
        filepath = os.path.join(self.output_dir, filename)

        doc = SimpleDocTemplate(
            filepath,
            pagesize=landscape(A4),
            rightMargin=15 * mm,
            leftMargin=15 * mm,
            topMargin=20 * mm,
            bottomMargin=15 * mm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Title"],
            fontSize=18,
            spaceAfter=20,
        )

        elements = []
        elements.append(Paragraph(f"月度考勤统计报告 - {year_month}", title_style))
        elements.append(Spacer(1, 10))

        table_data = [[
            "工号", "姓名", "部门", "应出勤", "实出勤",
            "迟到", "早退", "缺勤", "缺卡",
            "总工时", "平日加班", "周末加班", "节假日加班",
            "扣减", "补贴", "净调整", "达标率",
        ]]

        for summary in summaries:
            emp = await session.get(Employee, summary.employee_id)
            dept = await session.get(Department, summary.department_id)
            table_data.append([
                emp.employee_no if emp else "",
                emp.name if emp else "",
                dept.name if dept else "",
                str(summary.total_work_days),
                str(summary.actual_work_days),
                str(summary.late_count),
                str(summary.early_leave_count),
                str(summary.absent_count),
                str(summary.missing_punch_count),
                f"{float(summary.total_work_hours):.1f}",
                f"{float(summary.overtime_hours_weekday):.1f}",
                f"{float(summary.overtime_hours_weekend):.1f}",
                f"{float(summary.overtime_hours_holiday):.1f}",
                f"{float(summary.salary_deduction):.0f}",
                f"{float(summary.overtime_allowance):.0f}",
                f"{float(summary.net_adjustment):.0f}",
                f"{float(summary.compliance_rate) * 100:.1f}%",
            ])

        col_widths = [35, 40, 50, 30, 30, 25, 25, 25, 25, 35, 35, 35, 40, 35, 35, 35, 35]
        table = Table(table_data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("FONTSIZE", (0, 1), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ]))

        elements.append(table)

        elements.append(Spacer(1, 20))
        gen_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        elements.append(Paragraph(f"生成时间: {gen_time}", styles["Normal"]))

        doc.build(elements)
        logger.info("PDF报告已生成: %s", filepath)
        return filepath


class DashboardGenerator:
    def __init__(self, config: SystemConfig = None):
        self.config = config or get_config()

    async def generate_daily_metrics(self, target_date: date, session: AsyncSession) -> List[DashboardMetric]:
        stmt = select(Department).where(Department.is_active == True)
        result = await session.execute(stmt)
        departments = list(result.scalars().all())

        metrics = []
        for dept in departments:
            stmt_emp = select(func.count()).where(
                and_(Employee.department_id == dept.id, Employee.is_active == True)
            )
            total_employees = (await session.execute(stmt_emp)).scalar() or 0

            stmt_present = select(func.count(func.distinct(PunchRecord.employee_id))).where(
                and_(
                    PunchRecord.punch_date == target_date,
                    PunchRecord.is_valid == True,
                    PunchRecord.employee_id.in_(
                        select(Employee.id).where(
                            and_(Employee.department_id == dept.id, Employee.is_active == True)
                        )
                    ),
                )
            )
            present_employees = (await session.execute(stmt_present)).scalar() or 0

            attendance_rate = Decimal(str(round(present_employees / max(total_employees, 1), 4)))

            stmt_no_punch = select(func.count()).where(
                and_(
                    Employee.department_id == dept.id,
                    Employee.is_active == True,
                    Employee.id.notin_(
                        select(func.distinct(PunchRecord.employee_id)).where(
                            and_(PunchRecord.punch_date == target_date, PunchRecord.is_valid == True)
                        )
                    ),
                )
            )
            no_punch_count = (await session.execute(stmt_no_punch)).scalar() or 0

            stmt_anomaly = select(
                AnomalyRecord.anomaly_type,
                func.count().label("cnt"),
            ).where(
                and_(
                    AnomalyRecord.anomaly_date == target_date,
                    AnomalyRecord.employee_id.in_(
                        select(Employee.id).where(
                            and_(Employee.department_id == dept.id, Employee.is_active == True)
                        )
                    ),
                    AnomalyRecord.status != AnomalyStatus.AUTO_CORRECTED,
                )
            ).group_by(AnomalyRecord.anomaly_type)
            anomaly_result = await session.execute(stmt_anomaly)
            anomaly_counts = {row.anomaly_type: row.cnt for row in anomaly_result}

            anomaly_count = sum(anomaly_counts.values())
            late_count = anomaly_counts.get(AnomalyType.LATE, 0)
            early_leave_count = anomaly_counts.get(AnomalyType.EARLY_LEAVE, 0)

            existing_stmt = select(DashboardMetric).where(
                and_(
                    DashboardMetric.metric_date == target_date,
                    DashboardMetric.department_id == dept.id,
                )
            )
            existing = (await session.execute(existing_stmt)).scalar_one_or_none()

            if existing:
                metric = existing
            else:
                metric = DashboardMetric(
                    metric_date=target_date,
                    department_id=dept.id,
                )
                session.add(metric)

            metric.total_employees = total_employees
            metric.present_employees = present_employees
            metric.attendance_rate = attendance_rate
            metric.no_punch_count = no_punch_count
            metric.anomaly_count = anomaly_count
            metric.late_count = late_count
            metric.early_leave_count = early_leave_count

            metrics.append(metric)

        if metrics:
            await session.flush()

        logger.info("看板数据生成完成 [%s]: %d 个部门", target_date, len(metrics))
        return metrics

    async def get_dashboard_data(self, target_date: date, session: AsyncSession) -> Dict:
        stmt = select(DashboardMetric).where(DashboardMetric.metric_date == target_date)
        result = await session.execute(stmt)
        metrics = list(result.scalars().all())

        if not metrics:
            metrics = await self.generate_daily_metrics(target_date, session)

        total_employees = sum(m.total_employees for m in metrics)
        total_present = sum(m.present_employees for m in metrics)
        overall_rate = round(total_present / max(total_employees, 1), 4)

        dept_data = []
        for m in metrics:
            dept = await session.get(Department, m.department_id)
            dept_data.append({
                "department_id": m.department_id,
                "department_name": dept.name if dept else "未知",
                "total_employees": m.total_employees,
                "present_employees": m.present_employees,
                "attendance_rate": f"{float(m.attendance_rate) * 100:.1f}%",
                "no_punch_count": m.no_punch_count,
                "anomaly_count": m.anomaly_count,
                "late_count": m.late_count,
                "early_leave_count": m.early_leave_count,
            })

        return {
            "date": target_date.isoformat(),
            "company_overview": {
                "total_employees": total_employees,
                "total_present": total_present,
                "overall_attendance_rate": f"{overall_rate * 100:.1f}%",
                "total_no_punch": sum(m.no_punch_count for m in metrics),
                "total_anomalies": sum(m.anomaly_count for m in metrics),
            },
            "department_details": dept_data,
        }
